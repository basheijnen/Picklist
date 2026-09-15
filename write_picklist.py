import openpyxl

GEBIED_ORDER = ["KOELING", "KAS", "KAMER", "POKON", "DOZEN"]


def _display_aantal(aantal):
    return int(aantal) if float(aantal).is_integer() else aantal


def write_picklist(totals, unknown, output_path, for_date):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    for gebied in GEBIED_ORDER:
        ws = workbook.create_sheet(gebied)
        ws.append([f"E-COMMERCE {gebied} PICKLIST"])
        ws.append(["Datum:", for_date.strftime("%d-%m-%Y")])
        ws.append([])
        ws.append(["Item", "Soort", "Aantal"])
        rows = sorted(totals.get(gebied, {}).items(), key=lambda kv: (kv[0][0], kv[0][1]))
        for (item, soort), aantal in rows:
            ws.append([item, soort, _display_aantal(aantal)])

    unknown_ws = workbook.create_sheet("Onbekende pakketten")
    unknown_ws.append(["Pakketnummer", "Aantal besteld"])
    for pakketnummer, aantal in sorted(unknown.items()):
        unknown_ws.append([pakketnummer, _display_aantal(aantal)])

    workbook.save(output_path)
