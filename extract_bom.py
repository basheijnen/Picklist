import re
import sys

import openpyxl

from bom import BomEntry, write_bom_csv

INVOER_SHEET = "1. Invoer pakketaantal"
AREA_SHEETS = {
    "2. KOELING PICKLIST": "KOELING",
    "3. KAS PICKLIST": "KAS",
    "4. KAMER PICKLIST": "KAMER",
    "5. POKON PICKLIST": "POKON",
}

_GROUP_RE = re.compile(r"\(([^()]*)\)\*(-?\d+(?:\.\d+)?)")
_STANDALONE_RE = re.compile(r"'1\. Invoer pakketaantal'!G(\d+)\*(-?\d+(?:\.\d+)?)")
_CELLREF_RE = re.compile(r"'1\. Invoer pakketaantal'!G(\d+)")


def parse_formula(formula):
    """Parse a picklist formula that sums (cellref[+cellref...])*multiplier
    terms referencing '1. Invoer pakketaantal'!G<row>, into a list of
    (row, multiplier) pairs. Raises ValueError for any formula shape this
    doesn't recognise, so an unexpected future formula fails loudly instead
    of silently producing a wrong BOM.
    """
    terms = []
    for group_text, multiplier in _GROUP_RE.findall(formula):
        for row_str in _CELLREF_RE.findall(group_text):
            terms.append((int(row_str), float(multiplier)))
    remainder = _GROUP_RE.sub("", formula)

    for row_str, multiplier in _STANDALONE_RE.findall(remainder):
        terms.append((int(row_str), float(multiplier)))
    remainder = _STANDALONE_RE.sub("", remainder)

    if _CELLREF_RE.search(remainder):
        raise ValueError(f"Kan formule niet ontleden: {formula!r}")
    clean_remainder = (
        remainder.strip().lstrip("=").replace("+", "").replace("\n", "").strip()
    )
    if clean_remainder:
        raise ValueError(f"Kan formule niet ontleden: {formula!r}")

    return terms


def build_row_to_pakket(invoer_ws, first_row=5):
    row_to_pakket = {}
    for row in range(first_row, invoer_ws.max_row + 1):
        value = invoer_ws[f"B{row}"].value
        if value is not None:
            row_to_pakket[row] = str(value)
    return row_to_pakket


def extract_area_bom(workbook, sheet_name, gebied, row_to_pakket, first_row=6):
    ws = workbook[sheet_name]
    entries = []
    for row in range(first_row, ws.max_row + 1):
        item = ws[f"A{row}"].value
        soort = ws[f"B{row}"].value
        formula = ws[f"C{row}"].value
        if not isinstance(formula, str) or "Invoer" not in formula:
            continue
        if item is None and soort is None:
            continue
        for pakket_row, multiplier in parse_formula(formula):
            pakketnummer = row_to_pakket.get(pakket_row)
            if pakketnummer is None:
                raise ValueError(
                    f"{sheet_name} rij {row} verwijst naar Invoer-rij {pakket_row}, "
                    "die geen pakketnummer heeft"
                )
            entries.append(
                BomEntry(
                    pakketnummer=pakketnummer,
                    gebied=gebied,
                    item=str(item) if item is not None else "",
                    soort=str(soort) if soort is not None else "",
                    aantal_per_pakket=multiplier,
                )
            )
    return entries


def extract_dozen_bom(invoer_ws, first_row=5):
    entries = []
    for row in range(first_row, invoer_ws.max_row + 1):
        pakketnummer = invoer_ws[f"B{row}"].value
        doos_value = invoer_ws[f"H{row}"].value
        if pakketnummer is None or doos_value is None:
            continue
        for token in str(doos_value).split("+"):
            box = token.strip()
            if box:
                entries.append(
                    BomEntry(
                        pakketnummer=str(pakketnummer),
                        gebied="DOZEN",
                        item=box,
                        soort="",
                        aantal_per_pakket=1.0,
                    )
                )
    return entries


def extract_bom(xlsm_path):
    workbook = openpyxl.load_workbook(xlsm_path, data_only=False, keep_vba=True)
    invoer_ws = workbook[INVOER_SHEET]
    row_to_pakket = build_row_to_pakket(invoer_ws)

    entries = []
    for sheet_name, gebied in AREA_SHEETS.items():
        entries.extend(extract_area_bom(workbook, sheet_name, gebied, row_to_pakket))
    entries.extend(extract_dozen_bom(invoer_ws))
    return entries


if __name__ == "__main__":
    source_path = sys.argv[1] if len(sys.argv) > 1 else (
        "E-Commerce PICKLIST - IN PROGRESS.xlsm"
    )
    bom_entries = extract_bom(source_path)
    write_bom_csv(bom_entries, "bom.csv")
    print(f"{len(bom_entries)} BOM-regels geschreven naar bom.csv")
