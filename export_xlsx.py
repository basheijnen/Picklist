import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def maak_xlsx(rijen, bladnaam):
    """Eén werkblad met `rijen` (eerste rij = kolomkoppen) als .xlsx-bytes.
    Getallen blijven getallen, zodat je er in Excel meteen mee kunt rekenen.
    """
    wb = Workbook()
    ws = wb.active
    # Excel staat maximaal 31 tekens en geen []:*?/\ in een bladnaam toe.
    ws.title = "".join(teken for teken in str(bladnaam) if teken not in "[]:*?/\\")[:31] or "Blad1"
    for rij in rijen:
        ws.append(rij)
    for cel in ws[1]:
        cel.font = Font(bold=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for index, kolom in enumerate(ws.iter_cols(values_only=True), start=1):
        breedte = max((len(str(waarde)) for waarde in kolom if waarde is not None), default=8)
        ws.column_dimensions[get_column_letter(index)].width = min(max(breedte + 2, 10), 60)
    uitvoer = io.BytesIO()
    wb.save(uitvoer)
    return uitvoer.getvalue()
