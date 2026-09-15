import pytest

from extract_bom import parse_formula


def test_parse_formula_multiple_parenthesised_groups():
    formula = (
        "=('1. Invoer pakketaantal'!G27+'1. Invoer pakketaantal'!G28)*1"
        "+('1. Invoer pakketaantal'!G33+'1. Invoer pakketaantal'!G34)*3"
    )

    terms = parse_formula(formula)

    assert terms == [(27, 1.0), (28, 1.0), (33, 3.0), (34, 3.0)]


def test_parse_formula_standalone_terms_across_lines():
    formula = (
        "='1. Invoer pakketaantal'!G58*1\n"
        "+'1. Invoer pakketaantal'!G60*1"
    )

    terms = parse_formula(formula)

    assert terms == [(58, 1.0), (60, 1.0)]


def test_parse_formula_rejects_unsupported_shapes():
    with pytest.raises(ValueError):
        parse_formula("=SUM('1. Invoer pakketaantal'!G1:G5)")


def test_parse_formula_rejects_unrecognised_extra_term_inside_group():
    with pytest.raises(ValueError):
        parse_formula("=('1. Invoer pakketaantal'!G5+2)*3")


def test_parse_formula_rejects_unrecognised_column_inside_group():
    with pytest.raises(ValueError):
        parse_formula(
            "=('1. Invoer pakketaantal'!G5+'1. Invoer pakketaantal'!H6)*3"
        )


def test_parse_formula_rejects_unrecognised_other_sheet_ref_inside_group():
    with pytest.raises(ValueError):
        parse_formula("=('1. Invoer pakketaantal'!G5+'Blad2'!G9)*3")


import openpyxl as _openpyxl

from bom import BomEntry
from extract_bom import extract_dozen_bom


def _make_invoer_ws(rows):
    workbook = _openpyxl.Workbook()
    ws = workbook.active
    ws.title = "1. Invoer pakketaantal"
    for row_index, (pakketnummer, doos) in enumerate(rows, start=5):
        ws[f"B{row_index}"] = pakketnummer
        ws[f"H{row_index}"] = doos
    return ws


def test_extract_dozen_bom_splits_combined_box_values():
    ws = _make_invoer_ws([("1.10p", "14 + 15"), ("1.1", 14)])

    entries = extract_dozen_bom(ws)

    assert entries == [
        BomEntry("1.10p", "DOZEN", "14", "", 1.0),
        BomEntry("1.10p", "DOZEN", "15", "", 1.0),
        BomEntry("1.1", "DOZEN", "14", "", 1.0),
    ]
