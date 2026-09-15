from datetime import date

import openpyxl
import pytest

from bom import BomEntry, write_bom_csv
from generate_picklist import main


def _make_bron(path):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    ws = workbook.create_sheet("Totaal alles")
    for _ in range(5):
        ws.append([])
    ws.append(["Pakketnummer", "Aantal"])
    ws.append(["1.3", 2])
    ws.append(["99.9", 7])
    workbook.save(path)


def test_main_writes_picklist_and_reports_unknown(tmp_path, capsys):
    base_order_dir = tmp_path / "orders"
    day_folder = base_order_dir / "2026" / "16-09-2026"
    day_folder.mkdir(parents=True)
    _make_bron(day_folder / "Bron.xlsm")

    bom_csv_path = tmp_path / "bom.csv"
    write_bom_csv([BomEntry("1.3", "KOELING", "Parade", "CL Pink", 1.0)], bom_csv_path)

    output_path = tmp_path / "Picklist.xlsx"

    exit_code = main(
        base_order_dir=base_order_dir,
        bom_csv_path=bom_csv_path,
        output_path=output_path,
        today=date(2026, 9, 16),
    )

    assert exit_code == 0
    assert output_path.exists()
    workbook = openpyxl.load_workbook(output_path)
    assert workbook["KOELING"]["A5"].value == "Parade"
    assert "99.9" in capsys.readouterr().out


def test_main_returns_error_code_when_bron_missing(tmp_path, capsys):
    base_order_dir = tmp_path / "orders"
    bom_csv_path = tmp_path / "bom.csv"
    write_bom_csv([], bom_csv_path)
    output_path = tmp_path / "Picklist.xlsx"

    exit_code = main(
        base_order_dir=base_order_dir,
        bom_csv_path=bom_csv_path,
        output_path=output_path,
        today=date(2026, 9, 16),
    )

    assert exit_code == 1
    assert not output_path.exists()
    assert "FOUT" in capsys.readouterr().out


def test_main_returns_error_code_when_output_path_cannot_be_written(
    tmp_path, capsys, monkeypatch
):
    base_order_dir = tmp_path / "orders"
    day_folder = base_order_dir / "2026" / "16-09-2026"
    day_folder.mkdir(parents=True)
    _make_bron(day_folder / "Bron.xlsm")

    bom_csv_path = tmp_path / "bom.csv"
    write_bom_csv([BomEntry("1.3", "KOELING", "Parade", "CL Pink", 1.0)], bom_csv_path)
    output_path = tmp_path / "Picklist.xlsx"

    def _locked(*args, **kwargs):
        raise PermissionError("[Errno 13] Permission denied: 'Picklist.xlsx'")

    # Simulates the daily real-world failure mode: Picklist.xlsx is still
    # open in Excel when workbook.save() tries to overwrite it.
    monkeypatch.setattr("generate_picklist.write_picklist", _locked)

    exit_code = main(
        base_order_dir=base_order_dir,
        bom_csv_path=bom_csv_path,
        output_path=output_path,
        today=date(2026, 9, 16),
    )

    assert exit_code == 1
    out = capsys.readouterr().out
    assert "FOUT" in out
    assert "Picklist.xlsx" in out
    assert "Excel" in out


def test_main_returns_error_code_for_unexpected_exception(tmp_path, capsys, monkeypatch):
    base_order_dir = tmp_path / "orders"
    day_folder = base_order_dir / "2026" / "16-09-2026"
    day_folder.mkdir(parents=True)
    _make_bron(day_folder / "Bron.xlsm")

    bom_csv_path = tmp_path / "bom.csv"
    write_bom_csv([BomEntry("1.3", "KOELING", "Parade", "CL Pink", 1.0)], bom_csv_path)
    output_path = tmp_path / "Picklist.xlsx"

    def _boom(*args, **kwargs):
        raise RuntimeError("iets onverwachts")

    monkeypatch.setattr("generate_picklist.write_picklist", _boom)

    exit_code = main(
        base_order_dir=base_order_dir,
        bom_csv_path=bom_csv_path,
        output_path=output_path,
        today=date(2026, 9, 16),
    )

    assert exit_code == 1
    out = capsys.readouterr().out
    assert "FOUT" in out
    assert "iets onverwachts" in out


def test_main_uses_source_csv_path_instead_of_bron_when_given(tmp_path, capsys):
    csv_path = tmp_path / "export.csv"
    csv_path.write_text(
        '"Ordernr. intern";"Package Number";"Status"\n'
        '"1";"1.3";"Printed"\n'
        '"2";"1.3";"Printed"\n',
        encoding="utf-8-sig",
    )

    bom_csv_path = tmp_path / "bom.csv"
    write_bom_csv([BomEntry("1.3", "KOELING", "Parade", "CL Pink", 1.0)], bom_csv_path)

    output_path = tmp_path / "Picklist.xlsx"

    exit_code = main(
        base_order_dir=tmp_path / "unused",
        bom_csv_path=bom_csv_path,
        output_path=output_path,
        today=date(2026, 9, 16),
        source_csv_path=csv_path,
    )

    assert exit_code == 0
    assert output_path.exists()
    workbook = openpyxl.load_workbook(output_path)
    assert workbook["KOELING"]["A5"].value == "Parade"
    assert workbook["KOELING"]["C5"].value == 2


def test_main_returns_error_code_for_malformed_csv_source(tmp_path, capsys):
    # CSV missing the required "Package Number" column
    csv_path = tmp_path / "malformed.csv"
    csv_path.write_text(
        '"Ordernr. intern";"Status"\n'
        '"1";"Printed"\n',
        encoding="utf-8-sig",
    )

    bom_csv_path = tmp_path / "bom.csv"
    write_bom_csv([], bom_csv_path)
    output_path = tmp_path / "Picklist.xlsx"

    exit_code = main(
        base_order_dir=tmp_path / "unused",
        bom_csv_path=bom_csv_path,
        output_path=output_path,
        today=date(2026, 9, 16),
        source_csv_path=csv_path,
    )

    assert exit_code == 1
    assert not output_path.exists()
    out = capsys.readouterr().out
    assert "FOUT" in out
    assert "kan orders niet inlezen" in out
