from datetime import date

import openpyxl

from write_picklist import write_picklist


def test_write_picklist_creates_one_sheet_per_gebied_plus_unknown(tmp_path):
    totals = {
        "KOELING": {("Parade", "CL Pink"): 14.0},
        "DOZEN": {("9", ""): 3.0},
    }
    unknown = {"99.9": 7}
    output_path = tmp_path / "Picklist.xlsx"

    write_picklist(totals, unknown, output_path, date(2026, 9, 16))

    workbook = openpyxl.load_workbook(output_path)
    assert workbook.sheetnames == [
        "KOELING",
        "KAS",
        "KAMER",
        "POKON",
        "DOZEN",
        "Onbekende pakketten",
    ]

    koeling = workbook["KOELING"]
    assert koeling["B2"].value == "16-09-2026"
    assert [cell.value for cell in koeling[5]] == ["Parade", "CL Pink", 14]

    kas = workbook["KAS"]
    assert kas.max_row == 4  # header rows only, no items ordered

    onbekend = workbook["Onbekende pakketten"]
    assert [cell.value for cell in onbekend[2]] == ["99.9", 7]


def test_write_picklist_warns_about_unknown_gebied(tmp_path, capsys):
    totals = {
        "KOELING": {("Parade", "CL Pink"): 14.0},
        "Koeling": {("Typo", "Fout"): 5.0},
    }
    unknown = {}
    output_path = tmp_path / "Picklist.xlsx"

    write_picklist(totals, unknown, output_path, date(2026, 9, 16))

    captured = capsys.readouterr()
    assert "LET OP" in captured.out
    assert "Koeling" in captured.out
