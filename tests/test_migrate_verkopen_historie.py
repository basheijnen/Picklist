import datetime

import openpyxl
import pytest

from sales import load_verkoop_csv
from tools.migrate_verkopen_historie import migreer, read_kanaal_historie


def _make_workbook(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    amazon = wb.create_sheet("Amazon")
    amazon["B6"] = "Amazon"  # real workbooks give this the same value as the tab name
    amazon["D8"] = "9.1"
    amazon["E8"] = "70.12"
    amazon["A9"] = datetime.date(2026, 7, 1)
    amazon["D9"] = 5
    amazon["E9"] = 0  # zero cells must not become rows
    amazon["A10"] = "Totaal"

    # Real workbooks give B6 a Dutch display label that differs from the tab
    # name (e.g. tab "Groupon DE" has B6 "Groupon Duitsland") — the migration
    # must key off the tab name, since that's what matches the live app's
    # channel table, not this display label.
    # Real sheets can have a blank row 9 (a leftover week-52 subtotal row
    # with no date) before the actual dates start at row 10 — the scan must
    # skip past it rather than stopping there.
    groupon = wb.create_sheet("Groupon DE")
    groupon["B6"] = "Groupon Duitsland"
    groupon["D8"] = "9.8"
    groupon["B9"] = 52
    groupon["A10"] = datetime.date(2026, 7, 1)
    groupon["D10"] = 3
    groupon["A11"] = "Totaal"

    summary = wb.create_sheet("Verkopen per pakketnummer")
    summary["A1"] = "moet worden overgeslagen"

    wb.save(path)


def test_read_kanaal_historie_skips_summary_sheets_and_zero_cells(tmp_path):
    workbook_path = tmp_path / "historie.xlsm"
    _make_workbook(workbook_path)

    orders = read_kanaal_historie(workbook_path)

    assert len(orders) == 2
    by_kanaal = {order.kanaal: order for order in orders}
    amazon_order = by_kanaal["Amazon"]
    assert amazon_order.datum == "2026-07-01"
    assert amazon_order.pakketnummer == "9.1"
    assert amazon_order.aantal == 5.0
    assert amazon_order.ordernummer == "migratie-Amazon-2026-07-01-9.1"

    # Tab name "Groupon DE" must win over B6's "Groupon Duitsland".
    groupon_order = by_kanaal["Groupon DE"]
    assert groupon_order.pakketnummer == "9.8"
    assert groupon_order.aantal == 3.0


def test_read_kanaal_historie_splits_pokon_from_pakket(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bol = wb.create_sheet("Bol.com")
    bol["B6"] = "Bol.com"
    bol["D8"] = "9.1"
    bol["E8"] = "9.1P"  # package 9.1 sold with a box of Pokon added
    bol["F8"] = "P004"  # fixed internal code series, unrelated — ignored entirely
    bol["A9"] = datetime.date(2026, 7, 6)
    bol["D9"] = 30
    bol["E9"] = 1
    bol["F9"] = 1
    bol["A10"] = "Totaal"
    workbook_path = tmp_path / "historie.xlsm"
    wb.save(workbook_path)

    orders = read_kanaal_historie(workbook_path)

    by_pakketnummer = {}
    for order in orders:
        by_pakketnummer.setdefault(order.pakketnummer, 0)
        by_pakketnummer[order.pakketnummer] += order.aantal
    # 30 plain "9.1" + 1 "9.1P" (counts toward "9.1" itself) = 31.
    assert by_pakketnummer["9.1"] == 31.0
    # Only "9.1P"'s Pokon box counts — the standalone "P004" is ignored.
    assert by_pakketnummer["Pokon"] == 1.0
    assert "9.1P" not in by_pakketnummer
    assert "P004" not in by_pakketnummer


def test_migreer_writes_orders_and_is_idempotent(tmp_path):
    workbook_path = tmp_path / "historie.xlsm"
    _make_workbook(workbook_path)
    verkoop_csv_path = tmp_path / "verkoop_orders.csv"

    toegevoegd, overgeslagen, kanalen = migreer(workbook_path, verkoop_csv_path)
    assert toegevoegd == 2
    assert overgeslagen == 0
    assert kanalen == {"Amazon", "Groupon DE"}
    assert len(load_verkoop_csv(verkoop_csv_path)) == 2

    # Running it again must not duplicate rows.
    toegevoegd, overgeslagen, kanalen = migreer(workbook_path, verkoop_csv_path)
    assert toegevoegd == 0
    assert overgeslagen == 2
    assert kanalen == {"Amazon", "Groupon DE"}
    assert len(load_verkoop_csv(verkoop_csv_path)) == 2
