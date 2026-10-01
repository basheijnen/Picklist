import io
import json
import threading
import urllib.request
from http.server import ThreadingHTTPServer

import openpyxl

import webapp_server
from export_xlsx import maak_xlsx


def test_maak_xlsx_zet_rijen_in_een_werkblad_met_getallen_als_getal():
    inhoud = maak_xlsx([["Plant", "Potmaat", "Aantal verkocht"], ["Ficus Carica - P9", "P9", 419], ["Lavendel", "P17 / C2", 2.5]], "Per plant")

    ws = openpyxl.load_workbook(io.BytesIO(inhoud)).active
    assert ws.title == "Per plant"
    assert [[cel.value for cel in rij] for rij in ws.iter_rows()] == [
        ["Plant", "Potmaat", "Aantal verkocht"],
        ["Ficus Carica - P9", "P9", 419],
        ["Lavendel", "P17 / C2", 2.5],
    ]
    assert ws["A1"].font.bold
    assert ws.freeze_panes == "A2"


def test_maak_xlsx_kort_een_te_lange_bladnaam_in():
    inhoud = maak_xlsx([["a"]], "Een hele lange naam voor een werkblad")

    assert len(openpyxl.load_workbook(io.BytesIO(inhoud)).active.title) <= 31


def test_server_geeft_xlsx_terug_via_http_post():
    server = ThreadingHTTPServer(("127.0.0.1", 0), webapp_server.PicklistRequestHandler)
    port = server.server_address[1]
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        request = urllib.request.Request(
            f"http://127.0.0.1:{port}/api/export-xlsx",
            data=json.dumps({"bladnaam": "Verkopen", "rijen": [["Pakket", "Aantal"], ["2.1", 3]]}).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=5) as response:
            content_type = response.headers["Content-Type"]
            inhoud = response.read()
    finally:
        server.shutdown()
        server.server_close()

    assert content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ws = openpyxl.load_workbook(io.BytesIO(inhoud)).active
    assert ws["A2"].value == "2.1"
    assert ws["B2"].value == 3
