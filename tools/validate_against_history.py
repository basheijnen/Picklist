"""Recompute picklist totals from bom.csv against a historical, already-filled
picklist workbook, and compare against that workbook's own cached formula
results. Run this after any change to bom.csv (e.g. adding a new product) to
confirm the BOM still matches how the spreadsheet actually calculates.

Usage: python tools/validate_against_history.py <historical_xlsm_path>
"""

import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from bom import load_bom_csv
from calculate import calculate_totals

AREA_SHEETS = {
    "2. KOELING PICKLIST": "KOELING",
    "3. KAS PICKLIST": "KAS",
    "4. KAMER PICKLIST": "KAMER",
    "5. POKON PICKLIST": "POKON",
}


def read_historical_aantallen(invoer_ws):
    aantallen = {}
    for row in range(5, invoer_ws.max_row + 1):
        pakketnummer = invoer_ws[f"B{row}"].value
        aantal = invoer_ws[f"G{row}"].value
        if pakketnummer is not None:
            aantallen[str(pakketnummer)] = aantal if isinstance(aantal, (int, float)) else 0
    return aantallen


def main(historical_path, bom_csv_path=Path(__file__).resolve().parent.parent / "bom.csv"):
    # Load workbook twice: once for values, once for formulas
    workbook_values = openpyxl.load_workbook(historical_path, data_only=True, keep_vba=True)
    workbook_formulas = openpyxl.load_workbook(historical_path, data_only=False, keep_vba=True)

    aantallen = read_historical_aantallen(workbook_values["1. Invoer pakketaantal"])
    bom_entries = load_bom_csv(bom_csv_path)
    totals, unknown = calculate_totals(bom_entries, aantallen)

    # Report unknown pakketnummers if any exist
    if unknown:
        print(f"WARNING: {len(unknown)} unknown pakketnummers found in historical data:")
        for pakketnummer, aantal in sorted(unknown.items()):
            print(f"  {pakketnummer}: {aantal}")

    checked = matches = mismatches = 0
    for sheet_name, gebied in AREA_SHEETS.items():
        ws_values = workbook_values[sheet_name]
        ws_formulas = workbook_formulas[sheet_name]
        for row in range(6, ws_values.max_row + 1):
            # Skip rows where C{row} formula doesn't contain "Invoer" (like extract_bom.py does)
            formula = ws_formulas[f"C{row}"].value
            if not isinstance(formula, str) or "Invoer" not in formula:
                continue

            item = ws_values[f"A{row}"].value
            soort = ws_values[f"B{row}"].value
            cached = ws_values[f"C{row}"].value
            if item is None and soort is None:
                continue
            cached = cached if isinstance(cached, (int, float)) else 0
            predicted = totals.get(gebied, {}).get((str(item), str(soort) if soort else ""), 0)
            checked += 1
            if abs(predicted - cached) > 1e-9:
                mismatches += 1
                print(f"MISMATCH {sheet_name} row {row} ({item}/{soort}): "
                      f"predicted={predicted} cached={cached}")
            else:
                matches += 1

    print(f"checked={checked} matches={matches} mismatches={mismatches}")
    return 0 if mismatches == 0 else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
