"""One-time migration: backfill verkoop_orders.csv from the existing
'Totaaloverzicht verkochte pakketten <seizoen>.xlsm' sales-tracking
workbook, so the webapp dashboard starts with the season's real history
instead of an empty slate. Safe to re-run — dedupes on ordernummer.
"""

import datetime
import re
import sys
from pathlib import Path

import openpyxl

PROJECT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_DIR))

from sales import VerkoopOrder, load_verkoop_csv, merge_new_orders, write_verkoop_csv

VERKOOP_CSV_PATH = PROJECT_DIR / "verkoop_orders.csv"

SUMMARY_SHEETS = {
    "verkopen per week",
    "verkopen per dag",
    "verkopen per pakketnummer",
    "cbs periode en land",
}


PUUR_POKON = re.compile(r"^[Pp]\d+$")
MET_POKON = re.compile(r"^(\d+\.\d+)[Pp]$")


def _split_pokon(pakketnummer):
    """A pakketnummer ending in "p"/"P" is that package with a box of Pokon
    added — the base package still counts as itself, and the Pokon box
    counts separately. A bare "P<nummer>" (e.g. "P004") is this workbook's
    own column series for standalone Pokon sales (one column per Pokon
    variant — "P001" = Tuinmest, "P004" = Mediterrane mest, etc. — not tied
    to any specific plant package), so it counts as pure Pokon too, with no
    underlying basis package. Returns (basis_pakketnummer_or_None, is_pokon).
    """
    if PUUR_POKON.fullmatch(pakketnummer):
        return None, True
    match = MET_POKON.fullmatch(pakketnummer)
    if match:
        return match.group(1), True
    return pakketnummer, False


def _laatste_pakket_kolom(ws):
    kolom = 4  # D
    while ws.cell(row=8, column=kolom).value not in (None, ""):
        kolom += 1
    return kolom - 1


def read_kanaal_historie(workbook_path):
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, keep_vba=True)
    orders = []
    for ws in workbook.worksheets:
        if ws.title.strip().lower() in SUMMARY_SHEETS:
            continue
        # The tab name, not B6, is the canonical channel identity: B6 holds a
        # Dutch display label ("Groupon Duitsland", "TMG") that doesn't match
        # what the live app's channel table expects ("Groupon DE", "Mediahuis")
        # — confirmed by comparing every tab name against its own B6 value.
        kanaal = ws.title.strip()
        if not kanaal:
            continue
        laatste_kolom = _laatste_pakket_kolom(ws)
        if laatste_kolom < 4:
            continue
        pakketnummers = {
            kolom: str(ws.cell(row=8, column=kolom).value).strip()
            for kolom in range(4, laatste_kolom + 1)
        }
        rij = 9
        laatste_rij = ws.max_row
        while rij <= laatste_rij:
            datum_waarde = ws.cell(row=rij, column=1).value
            if isinstance(datum_waarde, str) and datum_waarde.strip().lower() == "totaal":
                break
            if not isinstance(datum_waarde, (datetime.date, datetime.datetime)):
                # A blank (or otherwise non-date) row before the real dates
                # start does happen in real sheets (e.g. a leftover week-52
                # subtotal row with no date) — skip it rather than stopping,
                # and rely on the "Totaal" sentinel (or ws.max_row) to end.
                rij += 1
                continue
            datum = datum_waarde.strftime("%Y-%m-%d")
            for kolom, pakketnummer in pakketnummers.items():
                if not pakketnummer:
                    continue
                waarde = ws.cell(row=rij, column=kolom).value
                if not waarde:
                    continue
                try:
                    waarde = float(waarde)
                except (TypeError, ValueError):
                    print(
                        f"Waarschuwing: niet-numerieke waarde {waarde!r} genegeerd "
                        f"({ws.title}, {datum}, pakket {pakketnummer})"
                    )
                    continue
                basis, pokon = _split_pokon(pakketnummer)
                if basis:
                    orders.append(
                        VerkoopOrder(
                            ordernummer=f"migratie-{ws.title}-{datum}-{pakketnummer}",
                            datum=datum,
                            kanaal=kanaal,
                            pakketnummer=basis,
                            aantal=waarde,
                        )
                    )
                if pokon:
                    orders.append(
                        VerkoopOrder(
                            ordernummer=f"migratie-{ws.title}-{datum}-{pakketnummer}-pokon",
                            datum=datum,
                            kanaal=kanaal,
                            pakketnummer="Pokon",
                            aantal=waarde,
                        )
                    )
            rij += 1
    return orders


def migreer(workbook_path, verkoop_csv_path=VERKOOP_CSV_PATH):
    nieuwe_orders = read_kanaal_historie(workbook_path)
    bestaande_orders = load_verkoop_csv(verkoop_csv_path)
    alle_orders, toegevoegd, overgeslagen = merge_new_orders(bestaande_orders, nieuwe_orders)
    write_verkoop_csv(alle_orders, verkoop_csv_path)
    kanalen = {order.kanaal for order in nieuwe_orders}
    return toegevoegd, overgeslagen, kanalen


def main():
    if len(sys.argv) != 2:
        print("Gebruik: python tools/migrate_verkopen_historie.py <pad-naar-totaaloverzicht.xlsm>")
        return 1
    toegevoegd, overgeslagen, kanalen = migreer(Path(sys.argv[1]))
    print(f"{toegevoegd} historische orders toegevoegd, {overgeslagen} overgeslagen (al aanwezig).")
    print(f"Kanalen gevonden in het Excel-bestand: {', '.join(sorted(kanalen))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
