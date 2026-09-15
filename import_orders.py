from pathlib import Path

import openpyxl

TOTAAL_ALLES_SHEET = "Totaal alles"


def find_bron_file(base_dir, for_date):
    folder = Path(base_dir) / str(for_date.year) / for_date.strftime("%d-%m-%Y")
    bron_path = folder / "Bron.xlsm"
    if not bron_path.exists():
        raise FileNotFoundError(
            f"Geen Bron.xlsm gevonden voor {for_date.strftime('%d-%m-%Y')} "
            f"(verwacht: {bron_path})"
        )
    return bron_path


def read_totaal_alles(bron_path, first_row=7):
    workbook = openpyxl.load_workbook(bron_path, data_only=True, keep_vba=True)
    ws = workbook[TOTAAL_ALLES_SHEET]
    aantallen = {}
    for row in range(first_row, ws.max_row + 1):
        pakketnummer = ws[f"A{row}"].value
        aantal = ws[f"B{row}"].value
        if pakketnummer is None:
            continue
        pakketnummer = str(pakketnummer).strip()
        if pakketnummer.lower() == "totaal":
            continue
        aantallen[pakketnummer] = aantal if isinstance(aantal, (int, float)) else 0
    return aantallen
